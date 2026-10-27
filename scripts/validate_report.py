"""Check shared manuscript provenance, results, native equations and links."""
from pathlib import Path
import csv,hashlib,json,re
from docx import Document
root=Path(__file__).resolve().parents[1]
source=(root/'README.md').read_bytes();provenance=json.loads((root/'results/word_report_software.json').read_text())
assert provenance['source_sha256']==hashlib.sha256(source).hexdigest(),'Rebuild Word after changing README'
d=Document(root/'report/Multiancestry_GWAS_Methodology_and_Results.docx')
assert d.core_properties.author=='Vusumuzi Mabasa'
assert len(d.inline_shapes)==7 and len(d.tables)==6
math_count=len(d.element.xpath('.//m:oMath'));assert math_count>=25
texts=[p.text for p in d.paragraphs]
for heading in ['Abstract','Introduction','Methods','Results','Discussion','Conclusion','Reproducibility','References']:assert heading in texts,heading
assert any('This is a simulation study' in t for t in texts)
assert any('every phenotype is simulated' in t for t in texts)
for i,p in enumerate(d.paragraphs):
    if p._p.xpath('.//w:drawing'):assert len(d.paragraphs[i-1].text)>100,'Interpretation missing before figure'
with (root/'results/revision/validation_summary.tsv').open() as file:rows=list(csv.DictReader(file,delimiter='\t'))
lookup={(r['scenario'],r['ld'],r['name']):r for r in rows}
assert [len(t.rows) for t in d.tables]==[5,5,9,6,7,7]
for rr in d.tables[0].rows[1:]:
    a=[c.text for c in rr.cells];r=lookup[('null','matched',a[0])]
    assert a[2]==f"{float(r['any_false_positive']):.3f}";assert a[3]==f"{float(r['any_false_positive_mcse']):.3f}"
for rr in d.tables[2].rows[1:]:
    a=[c.text for c in rr.cells];scenario,name=a[0].split(' ',1);r=lookup[(scenario.lower(),'matched',name)]
    assert a[1:]==[f"{float(r[k]):.3f}" for k in ['recall','ld_recall_african','fdp','rmse_african']]
with (root/'results/revision/inla/comparison_summary.tsv').open() as file:inla_rows=list(csv.DictReader(file,delimiter='\t'))
inla_lookup={(r['method'],r['scenario']):r for r in inla_rows if r['ld']=='matched'}
for rr in d.tables[5].rows[1:]:
    a=[c.text for c in rr.cells];shared=inla_lookup[(a[0],'shared')];hetero=inla_lookup[(a[0],'heterogeneous')]
    expected=[f"{float(shared['rmse_african']):.3f}"]+[f"{float(hetero[k]):.3f}" if float(hetero[k])==float(hetero[k]) else '—' for k in ['rmse_african','specific_estimate','coverage_causal']]
    assert a[1:]==expected
links=[r.target_ref for r in d.part.rels.values() if r.reltype.endswith('/hyperlink')]
assert all(x.startswith(('https://','http://','#')) for x in links)
for dest in re.findall(r'\]\(([^)]+)\)',source.decode()):
    if not dest.startswith(('https://','#')):assert (root/dest).exists(),dest
assert (root/'workflow.txt').read_bytes()==source
for doi in ['10.1038/s41588-026-02799-1','10.1214/16-STS576']:assert any(doi in t for t in texts)
assert not re.search(r'@[A-Z]+@',source.decode())
audit=dict(author='Vusumuzi Mabasa',source_sha256=hashlib.sha256(source).hexdigest(),native_math_objects=math_count,figures=7,tables=6,numerical_tables_match=True,interpretations_before_figures=True,simulation_label=True,portable_links=True,visual_review='See results/report_validation.json for separate rendered-page inspection')
(root/'results/revision/manuscript_validation.json').write_text(json.dumps(audit,indent=2)+'\n',newline='\n');print(json.dumps(audit,indent=2))
